#!/usr/bin/env python3
"""표들을 엑셀 한 권으로 묶는다 — 팀이 바로 열어 쓰는 형태.

csv 는 엑셀에서 한글이 깨지기 쉽고 시트를 오가며 보기 어렵다. 한 권으로 묶고
머리줄 고정·자동 필터·자릿수 구분을 넣어 둔다.

나가는 곳: data/market/병의원_상권자료.xlsx
"""
import csv, pathlib, datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

SRC = pathlib.Path("data/market")
OUT = SRC / "병의원_상권자료.xlsx"
FONT = "맑은 고딕"          # 한글 문서의 기본 글꼴. 없는 기기에서는 시스템 고딕으로 떨어진다.
INK = "1C1E54"              # 베놈 남색
ACCENT = "533AFD"
DASH = "—"

SHEETS = [
    ("시군구", "market_sggu.csv", "시군구 한 줄 — 인구·면적·병의원·종별·밀집도·나이대·미용 겸업"),
    ("지역x진료과목", "market_subject.csv", "시군구 × 진료과목 — 내건 곳 / 전문의 있는 곳 / 전문의 없이"),
    ("지도분류", "market_kakao.csv", "카카오 지도가 그 병원을 무엇으로 분류했나"),
    ("업종별상권", "market_sangga.csv",
     "시군구 한 줄 — 업종별 가게 수와 «이 동네는 어떤 동네인가»"),
    ("행정동", "market_dong.csv", "행정동 한 줄 — 인구·남녀·나이대·면적"),
    ("검색수요", "naver_datalab_subjects.csv", "과목·시술별 검색 수요(피부과=100)"),
    ("사장님시트", "population_sggu.csv", "처음 주신 인구 시트를 옮긴 것"),
]


def as_value(s: str):
    """숫자는 숫자로, 못 잰 값 «—» 는 그대로 둔다(엑셀에서 맨 아래로 정렬된다)."""
    s = (s or "").strip()
    if s in ("", DASH):
        return DASH if s == DASH else None
    try:
        return int(s)
    except ValueError:
        pass
    try:
        return float(s)
    except ValueError:
        return s


def style_sheet(ws, ncols: int, widths: dict) -> None:
    head = Font(name=FONT, size=10, bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor=INK)
    thin = Side(style="thin", color="D8DEE9")
    for c in range(1, ncols + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = head; cell.fill = fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=thin)
        ws.column_dimensions[get_column_letter(c)].width = widths.get(c, 12)
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(ncols)}{ws.max_row}"


def main() -> int:
    wb = Workbook()
    wb.remove(wb.active)
    counts = {}

    for name, fname, _desc in SHEETS:
        path = SRC / fname
        if not path.exists():
            continue
        rows = list(csv.reader(path.open(encoding="utf-8-sig")))
        ws = wb.create_sheet(name)
        header = rows[0]
        ws.append(header)
        body = Font(name=FONT, size=10)
        for r in rows[1:]:
            ws.append([as_value(v) for v in r])
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.font = body
                if isinstance(cell.value, (int, float)):
                    cell.number_format = "#,##0" if isinstance(cell.value, int) else "#,##0.00"
        widths = {1: 9, 2: 16, 3: 16}
        style_sheet(ws, len(header), widths)
        counts[name] = ws.max_row - 1

    # ── 표지 ──
    cover = wb.create_sheet("읽는 법", 0)
    big = Font(name=FONT, size=18, bold=True, color=INK)
    sub = Font(name=FONT, size=10, color="64748D")
    h = Font(name=FONT, size=11, bold=True, color=ACCENT)
    b = Font(name=FONT, size=10)
    cover["A1"] = "지역별 병의원 상권 자료"; cover["A1"].font = big
    cover["A2"] = f"만든 날 {datetime.date.today():%Y-%m-%d} · 베놈 병원마케팅"; cover["A2"].font = sub

    r = 4
    cover[f"A{r}"] = "무엇이 들어 있나"; cover[f"A{r}"].font = h; r += 1
    cover[f"A{r}"] = "시트"; cover[f"B{r}"] = "줄 수"; cover[f"C{r}"] = "한 줄이 뜻하는 것"
    for c in "ABC":
        cover[f"{c}{r}"].font = Font(name=FONT, size=10, bold=True)
    r += 1
    first_data = r
    for name, _f, desc in SHEETS:
        if name not in counts:
            continue
        cover[f"A{r}"] = name
        cover[f"B{r}"] = counts[name]
        cover[f"C{r}"] = desc
        for c in "ABC":
            cover[f"{c}{r}"].font = b
        cover[f"B{r}"].number_format = "#,##0"
        r += 1

    r += 1
    cover[f"A{r}"] = "전국 합계 (시군구 시트에서 바로 센다)"; cover[f"A{r}"].font = h; r += 1
    n = counts.get("시군구", 0)
    # 계산식이 아니라 «세어 적은 값»을 싣는다. 이 방 환경에서는 계산식이 맞는지
    # 확인할 도구(리브레오피스)가 돌지 않는다 — 확인 못 한 계산식을 싣느니 센 값을 싣는다.
    # 시트가 바뀌면 이 도구를 다시 돌리면 표지도 같이 다시 적힌다.
    head_sggu = [c.value for c in wb["시군구"][1]]
    idx = {nm: i for i, nm in enumerate(head_sggu)}
    body_rows = [[c.value for c in row] for row in wb["시군구"].iter_rows(min_row=2)]
    tot = lambda nm: sum(v for x in body_rows
                         if isinstance(v := x[idx[nm]], (int, float)))
    pop, hosp, docs = tot("인구"), tot("병의원_계"), tot("의사수")
    for label, value, fmt in [
        ("인구", pop, "#,##0"),
        ("병의원", hosp, "#,##0"),
        ("의사 수", docs, "#,##0"),
        ("인구 1만 명당 병의원", hosp / pop * 10000 if pop else None, "#,##0.0"),
    ]:
        cover[f"A{r}"] = label; cover[f"A{r}"].font = b
        cover[f"B{r}"] = value; cover[f"B{r}"].font = b; cover[f"B{r}"].number_format = fmt
        r += 1
    cover[f"A{r}"] = "위 넷은 아래 「시군구」 시트를 세어 적은 값입니다(만든 날 기준)."
    cover[f"A{r}"].font = sub; r += 1

    r += 1
    for line in [
        ("어디서 온 자료인가", True),
        ("병의원·진료과목 — 건강보험심사평가원 「전국 병의원 및 약국 현황」 2026년 6월분", False),
        ("    (공공누리 제1유형 · 출처표시 — 이 자료를 쓴 화면·문서에는 출처를 적습니다)", False),
        ("인구·남녀 — 행정안전부 주민등록 인구통계 2026-06-30 · 나이대 2026-08-31", False),
        ("면적 — 통계청 SGIS 행정동 경계 2026-07-01", False),
        ("지도 분류 — 카카오 로컬 2026-09-19 (전국 훑기)", False),
        ("업종별 상권 — 소상공인시장진흥공단 상가(상권)정보 2026년 6월 기준", False),
        ("검색 수요 — 네이버 데이터랩 2025-09 ~ 2026-08", False),
        ("", False),
        ("표를 읽을 때 꼭 알아야 할 것", True),
        ("① 「내건 곳」은 전문 병원 수가 아닙니다. 심평원은 한 병원이 내건 진료과목을 모두 적습니다.", False),
        ("    피부과를 내건 곳이 전국 17,826곳인데 피부과 전문의가 있는 곳은 1,735곳뿐입니다.", False),
        ("    마케팅에서는 「그 검색어로 부딪칠 병원 수」라 이 숫자가 오히려 쓸모 있고,", False),
        ("    「진짜 그 과 병원」을 보려면 「전문의 있는 곳」 칸을 봅니다.", False),
        ("② 「지도분류」는 한 병원에 하나만 붙습니다 — 환자가 그 병원을 무엇으로 보는가입니다.", False),
        ("③ 「업종별상권」은 그 동네에 어떤 가게가 몇 곳 있는지입니다. 같은 인구라도", False),
        ("    학원가·먹자골목·오피스는 오는 사람이 다릅니다. 「전국대비」가 100보다 크면", False),
        ("    전국 평균보다 그 업종이 많다는 뜻입니다.", False),
        ("④ «—» 는 못 잰 값입니다. 지어내지 않았습니다. 지금 남은 것은 세종 나이대뿐입니다", False),
        ("    (행정안전부 나이대 자료에 세종이 통째로 빠져 있습니다).", False),
        ("", False),
        ("눌러 보며 찾는 화면", True),
        ("https://claude.ai/artifact/9nJf5n2wD54uuTF3hwywJL", False),
    ]:
        cover[f"A{r}"] = line[0]
        cover[f"A{r}"].font = h if line[1] else b
        r += 1

    cover.column_dimensions["A"].width = 46
    cover.column_dimensions["B"].width = 14
    cover.column_dimensions["C"].width = 60
    cover.sheet_view.showGridLines = False

    wb.save(OUT)
    print(f"{OUT} ({OUT.stat().st_size/1024:.0f}KB)")
    for k, v in counts.items():
        print(f"  {k:<16}{v:>7,}줄")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
